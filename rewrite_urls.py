#!/usr/bin/env python3
"""Swap the placehold.co Logo URLs for the Cordova division logo URLs.

Change BASE below to your own repo/CDN path and re-run to repoint everything.
"""
import sys
from collections import Counter
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

SRC = "/root/.claude/uploads/95942f24-7803-591e-9af8-0834b0e806e3/f9fc1dcd-Cordova_Pharmaceuticals_Assets.xlsx"
DST = "/mnt/user-data/outputs/Cordova_Pharmaceuticals_Assets.xlsx"

BASE = "https://raw.githubusercontent.com/YOUR-GITHUB-ORG/cordova-assets/main/logos"

# old placehold.co URL  ->  (manufacturer, new file slug)
MAP = {
    "https://placehold.co/96x96/1B6AC9/FFFFFF/png?text=CBS":
        ("Cordova BioSolutions", "cordova-biosolutions"),
    "https://placehold.co/96x96/2F6F3E/FFFFFF/png?text=CNS":
        ("Cordova NeuroSciences", "cordova-neurosciences"),
    "https://placehold.co/96x96/C2334D/FFFFFF/png?text=CCL":
        ("Cordova Cardio Labs", "cordova-cardio-labs"),
    "https://placehold.co/96x96/6B3FA0/FFFFFF/png?text=CO":
        ("Cordova Oncology", "cordova-oncology"),
    "https://placehold.co/96x96/E0761B/FFFFFF/png?text=CCH":
        ("Cordova Consumer Health", "cordova-consumer-health"),
    "https://placehold.co/96x96/0E9488/FFFFFF/png?text=CR":
        ("Cordova Respiratory", "cordova-respiratory"),
}
NEW = {old: f"{BASE}/{slug}.png" for old, (_, slug) in MAP.items()}


def main():
    wb = openpyxl.load_workbook(SRC)
    changed = Counter()
    mismatch = []
    unknown = Counter()

    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]]
        if "Logo URL" not in hdr:
            continue
        ic = hdr.index("Logo URL") + 1
        mc = hdr.index("Manufacturer") + 1
        for r in range(2, ws.max_row + 1):
            old = ws.cell(r, ic).value
            if not old:
                continue
            if old not in NEW:
                unknown[old] += 1
                continue
            brand = ws.cell(r, mc).value
            if brand != MAP[old][0]:
                mismatch.append((ws.title, r, brand, MAP[old][0]))
            ws.cell(r, ic).value = NEW[old]
            changed[(ws.title, MAP[old][0])] += 1

    # reference sheet: one row per division, with a live IMAGE() preview
    if "Logo Reference" in wb.sheetnames:
        del wb["Logo Reference"]
    ref = wb.create_sheet("Logo Reference")
    head = ["Division", "Initials", "Brand Color", "Preview", "Logo URL"]
    ref.append(head)
    for c in range(1, len(head) + 1):
        cell = ref.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3552")
        cell.alignment = Alignment(vertical="center")
    for old, (brand, slug) in MAP.items():
        initials = old.rsplit("=", 1)[1]
        color = "#" + old.split("/")[4]
        ref.append([brand, initials, color, None, NEW[old]])
    for r in range(2, ref.max_row + 1):
        ref.cell(r, 4).value = f'=IMAGE(E{r},"Cordova division logo",0)'
        ref.row_dimensions[r].height = 56
    for col, w in zip("ABCDE", (28, 10, 14, 12, 74)):
        ref.column_dimensions[col].width = w
    ref.freeze_panes = "A2"

    wb.save(DST)

    print(f"base url : {BASE}\n")
    for (sheet, brand), n in sorted(changed.items()):
        print(f"  {sheet:8s} {brand:26s} {n:5d} rows")
    print(f"\ntotal rewritten : {sum(changed.values())}")
    print(f"brand mismatches: {len(mismatch)}")
    if mismatch:
        for m in mismatch[:10]:
            print("   ", m)
    if unknown:
        print("unrecognised URLs:")
        for u, n in unknown.items():
            print(f"    {n:5d}  {u}")
    print(f"\nsaved -> {DST}")
    return 1 if (mismatch or unknown) else 0


if __name__ == "__main__":
    sys.exit(main())
